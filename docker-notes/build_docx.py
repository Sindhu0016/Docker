"""Build a print-ready Word document from notes.html."""

from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor, Twips

NAVY = RGBColor(0x1A, 0x36, 0x54)
INK = RGBColor(0x1C, 0x24, 0x30)
SLATE = RGBColor(0x5C, 0x73, 0x90)
HEAD3 = RGBColor(0x24, 0x3B, 0x53)

HERE = Path(__file__).resolve().parent
HTML = HERE / "notes.html"
OUT = Path(r"c:\Users\dasar\OneDrive\Desktop\Dasari_Sindhu_Docker_Notes.docx")


def _rpr(run):
    rpr = run._r.get_or_add_rPr()
    return rpr


def shade_run(run, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _rpr(run).append(shd)


def shade_paragraph(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def left_border(paragraph, color="7AA0C4", size="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    pBdr.append(left)
    pPr.append(pBdr)


def keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:keepNext")
    pPr.append(el)


def page_break_before(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    el = OxmlElement("w:pageBreakBefore")
    pPr.append(el)


def set_run_font(run, name, size, color=None, bold=False, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.name = name
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rfonts = run._element.rPr.rFonts
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, "Calibri", 9, SLATE)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def set_cell_shading(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "D5DEE8")
        borders.append(el)
    tcPr.append(borders)


def set_cell_margins(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", "60"), ("left", "80"), ("bottom", "60"), ("right", "80")):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), val)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


class Builder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.doc = Document()
        self._setup()
        self.stack = []
        self.section = None
        self.cover = {}
        self.callout = False
        self.list_kinds = []
        self.capture = None
        self.table_rows = None
        self.current_row = None
        self.cell = None
        self.ignore = 0

    def _setup(self):
        section = self.doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(20)
        section.right_margin = Mm(20)
        section.top_margin = Mm(18)
        section.bottom_margin = Mm(18)
        section.header_distance = Mm(0)
        section.footer_distance = Mm(0)
        section.different_first_page_header_footer = False
        self._set_valign(section, "center")
        self._cover_header_footer(section)
        self.usable = Mm(210) - Mm(36)

        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.font.color.rgb = INK
        pf = normal.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.space_after = Pt(8)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.15

        for style_name, size, color, before, after in (
            ("Heading 1", 18, NAVY, 0, 6),
            ("Heading 2", 13, NAVY, 12, 4),
            ("Heading 3", 12, HEAD3, 10, 4),
        ):
            style = self.doc.styles[style_name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.color.rgb = color
            style.font.bold = True
            sp = style.paragraph_format
            sp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            sp.space_before = Pt(before)
            sp.space_after = Pt(after)
            sp.line_spacing = 1.1
            sp.keep_with_next = True

        core = self.doc.core_properties
        core.title = "Docker Study Notes"
        core.author = "Dasari Sindhu"
        core.subject = "Docker study guide"
        core.category = "Study notes"

    def _set_valign(self, section, value):
        sectPr = section._sectPr
        for child in list(sectPr):
            if child.tag == qn("w:vAlign"):
                sectPr.remove(child)
        if value:
            align = OxmlElement("w:vAlign")
            align.set(qn("w:val"), value)
            sectPr.append(align)

    def _rule(self, paragraph, edge, color, size="12"):
        pPr = paragraph._p.get_or_add_pPr()
        pBdr = pPr.find(qn("w:pBdr"))
        if pBdr is None:
            pBdr = OxmlElement("w:pBdr")
            pPr.append(pBdr)
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), size)
        line.set(qn("w:space"), "1")
        line.set(qn("w:color"), color)
        pBdr.append(line)

    def _cover_header_footer(self, section):
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hp.paragraph_format.space_before = Pt(0)
        hp.paragraph_format.space_after = Pt(0)
        hp.clear()

        footer = section.footer
        footer.is_linked_to_previous = False
        footer.paragraphs[0].clear()

    def _start_body(self):
        body = self.doc.add_section(WD_SECTION.NEW_PAGE)
        body.page_width = Mm(210)
        body.page_height = Mm(297)
        body.left_margin = Mm(18)
        body.right_margin = Mm(18)
        body.top_margin = Mm(20)
        body.bottom_margin = Mm(18)
        body.header_distance = Mm(8)
        body.footer_distance = Mm(8)
        body.different_first_page_header_footer = False
        self._set_valign(body, None)
        self.usable = body.page_width - body.left_margin - body.right_margin
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), "1")
        body._sectPr.append(pg)
        self._header_footer(body)
        # The section break sits in an empty paragraph under the title card.
        cover_sect = self.doc.sections[0]._sectPr
        paragraph = cover_sect.getparent().getparent()
        if paragraph.tag == qn("w:p"):
            pPr = paragraph.find(qn("w:pPr"))
            if pPr is None:
                pPr = OxmlElement("w:pPr")
                paragraph.insert(0, pPr)
            spacing = pPr.find(qn("w:spacing"))
            if spacing is None:
                spacing = OxmlElement("w:spacing")
                pPr.append(spacing)
            spacing.set(qn("w:before"), "0")
            spacing.set(qn("w:after"), "0")
            spacing.set(qn("w:line"), "20")
            spacing.set(qn("w:lineRule"), "exact")

    def _header_footer(self, section):
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(2)
        left = p.add_run("Docker Study Notes")
        set_run_font(left, "Calibri", 9, SLATE)
        tab = p.add_run("\t")
        set_run_font(tab, "Calibri", 9, SLATE)
        right = p.add_run("Dasari Sindhu")
        set_run_font(right, "Calibri", 9, SLATE)
        pPr = p._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        tab_el = OxmlElement("w:tab")
        tab_el.set(qn("w:val"), "right")
        tab_el.set(qn("w:pos"), str(int(self.usable)))
        tabs.append(tab_el)
        pPr.append(tabs)
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "D5E0EC")
        pBdr.append(bottom)
        pPr.append(pBdr)

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.clear()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(2)
        pPr = fp._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "6")
        top.set(qn("w:space"), "1")
        top.set(qn("w:color"), "D5E0EC")
        pBdr.append(top)
        pPr.append(pBdr)
        add_page_number(fp)

    def classes(self, attrs):
        raw = ""
        for key, value in attrs:
            if key == "class":
                raw = value or ""
        return set(raw.split())

    def handle_starttag(self, tag, attrs):
        if tag in ("style", "script"):
            self.ignore += 1
            return
        if self.ignore:
            return
        classes = self.classes(attrs)
        self.stack.append((tag, classes))
        if tag == "section":
            if "cover" in classes:
                self.section = "cover"
            elif "toc" in classes:
                self._close_capture()
                self.section = "toc"
            return
        if tag == "div" and ("callout" in classes or "note" in classes):
            self.callout = True
            return
        if tag in ("ul", "ol"):
            self.list_kinds.append(tag)
            return
        if tag == "table":
            self._close_capture()
            self.table_rows = []
            return
        if tag == "tr":
            self.current_row = []
            return
        if tag in ("th", "td"):
            self.cell = {"header": tag == "th", "runs": []}
            return
        if tag in ("p", "h1", "h2", "h3", "li", "pre"):
            self._close_capture()
            self.capture = {
                "tag": tag,
                "classes": classes,
                "runs": [],
                "callout": self.callout,
                "section": self.section,
                "list": self.list_kinds[-1] if self.list_kinds else None,
            }
            return
        if self.capture is not None and tag == "br":
            self.capture["runs"].append(("\n", set()))

    def handle_endtag(self, tag):
        if tag in ("style", "script"):
            self.ignore = max(0, self.ignore - 1)
            return
        if self.ignore:
            return
        if tag in ("p", "h1", "h2", "h3", "li", "pre"):
            self._close_capture()
        elif tag in ("th", "td") and self.cell is not None:
            if self.current_row is not None:
                self.current_row.append(self.cell)
            self.cell = None
        elif tag == "tr" and self.current_row is not None:
            if self.table_rows is not None:
                self.table_rows.append(self.current_row)
            self.current_row = None
        elif tag == "table":
            self._render_table()
            self.table_rows = None
        elif tag == "div":
            self.callout = False
        elif tag in ("ul", "ol") and self.list_kinds:
            self.list_kinds.pop()
        elif tag == "section":
            self._close_capture()
            if self.section == "cover":
                self._paint_cover()
                self._start_body()
            self.section = None
        if self.stack and self.stack[-1][0] == tag:
            self.stack.pop()
        else:
            for i in range(len(self.stack) - 1, -1, -1):
                if self.stack[i][0] == tag:
                    self.stack = self.stack[:i]
                    break

    def handle_data(self, data):
        if self.ignore:
            return
        target = self.cell if self.cell is not None else self.capture
        if target is None:
            return
        pre = self.capture is not None and self.capture["tag"] == "pre" and self.cell is None
        if pre:
            text = data.replace("\r\n", "\n")
        else:
            text = re.sub(r"\s+", " ", data)
            if not text.strip() and not target["runs"]:
                return
        flags = set()
        for tag, _classes in self.stack:
            if tag in ("strong", "b", "th"):
                flags.add("bold")
            elif tag in ("em", "i"):
                flags.add("italic")
            elif tag == "code":
                flags.add("code")
        target["runs"].append((text, flags))

    def _flags_now(self):
        return set()

    def _close_capture(self):
        cap = self.capture
        self.capture = None
        if not cap:
            return
        text = "".join(part for part, _flags in cap["runs"])
        if cap["tag"] != "pre" and not text.strip():
            return
        if cap["tag"] == "pre" and not text.strip():
            return
        self._render(cap)

    def _bookmark(self, paragraph, name):
        self._bm_id = getattr(self, "_bm_id", 0) + 1
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self._bm_id))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self._bm_id))
        paragraph._p.insert(0, start)
        paragraph._p.append(end)

    def _add_internal_link(self, paragraph, text, anchor):
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), anchor)
        run = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        rStyle = OxmlElement("w:rStyle")
        rStyle.set(qn("w:val"), "Hyperlink")
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), "Calibri")
        rFonts.set(qn("w:hAnsi"), "Calibri")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "22")
        sizeCs = OxmlElement("w:szCs")
        sizeCs.set(qn("w:val"), "22")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "1A3654")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        for child in (rStyle, rFonts, size, sizeCs, color, underline):
            rPr.append(child)
        text_el = OxmlElement("w:t")
        text_el.set(qn("xml:space"), "preserve")
        text_el.text = text
        run.append(rPr)
        run.append(text_el)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    def _add_runs(self, paragraph, runs, *, pre=False, size=11, color=INK, name="Calibri"):
        if pre:
            raw = "".join(text for text, _flags in runs).replace("\r\n", "\n")
            raw = raw.strip("\n")
            lines = raw.split("\n")
            for index, line in enumerate(lines):
                run = paragraph.add_run(line if line else " ")
                set_run_font(run, "Consolas", 9, INK)
                if index < len(lines) - 1:
                    run.add_break()
            return
        for text, flags in runs:
            if text == "\n":
                paragraph.add_run().add_break()
                continue
            run = paragraph.add_run(text)
            font_name = "Consolas" if "code" in flags else name
            font_size = 10 if "code" in flags else size
            set_run_font(
                run,
                font_name,
                font_size,
                color,
                bold="bold" in flags,
                italic="italic" in flags,
            )
            if "code" in flags:
                shade_run(run, "F0F4F8")

    def _render(self, cap):
        tag = cap["tag"]
        classes = cap["classes"]
        section = cap["section"]

        if tag == "pre":
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Cm(0.15)
            shade_paragraph(p, "F6F8FA")
            left_border(p, "B7C9DB", "10")
            self._add_runs(p, cap["runs"], pre=True)
            return

        if section == "cover":
            text = "".join(part for part, _flags in cap["runs"]).strip()
            if tag == "h1":
                self.cover["title"] = text
            elif "subtitle" in classes:
                self.cover["subtitle"] = text
            elif "name" in classes:
                self.cover["name"] = text
            elif "meta" in classes:
                self.cover.setdefault("meta", []).append(text)
            return

        if "kicker" in classes:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(6)
            self._add_runs(p, cap["runs"], size=11, color=SLATE)
            for run in p.runs:
                run.bold = True
            return

        if "subtitle" in classes:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(14)
            self._add_runs(p, cap["runs"], size=14, color=RGBColor(0x3D, 0x51, 0x66))
            return

        if "name" in classes:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            self._add_runs(p, cap["runs"], size=16, color=INK)
            for run in p.runs:
                run.bold = True
            return

        if "meta" in classes:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            self._add_runs(p, cap["runs"], size=11, color=RGBColor(0x3D, 0x51, 0x66))
            return

        if "chapnum" in classes:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            page_break_before(p)
            keep_with_next(p)
            self._chapter_n = getattr(self, "_chapter_n", 0) + 1
            self._bookmark(p, f"chapter{self._chapter_n}")
            self._add_runs(p, cap["runs"], size=10, color=SLATE)
            for run in p.runs:
                run.bold = True
                rpr = _rpr(run)
                caps = OxmlElement("w:caps")
                rpr.append(caps)
            return

        if tag in ("h1", "h2", "h3"):
            style = {"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}[tag]
            p = self.doc.add_paragraph(style=style)
            self._add_runs(p, cap["runs"], size={"h1": 18, "h2": 13, "h3": 12}[tag], color=NAVY if tag != "h3" else HEAD3)
            for run in p.runs:
                run.bold = True
            return

        if section == "toc" and tag == "li":
            self._toc_n = getattr(self, "_toc_n", 0) + 1
            label = "".join(part for part, _flags in cap["runs"]).strip()
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            self._add_internal_link(p, f"{self._toc_n}.  {label}", f"chapter{self._toc_n}")
            return

        style_name = None
        if tag == "li":
            style_name = "List Number" if cap["list"] == "ol" else "List Bullet"
        p = self.doc.add_paragraph(style=style_name) if style_name else self.doc.add_paragraph()
        if tag == "li":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(0)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(8)
        if cap["callout"]:
            shade_paragraph(p, "F4F8FB")
            left_border(p, "7AA0C4", "16")
            p.paragraph_format.left_indent = Cm(0.2)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
        if "small" in classes:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            self._add_runs(p, cap["runs"], size=10, color=RGBColor(0x3D, 0x51, 0x66))
            return
        self._add_runs(p, cap["runs"])

    def _paint_cover(self):
        if self.doc.paragraphs and not self.doc.paragraphs[0].text.strip():
            self.doc.paragraphs[0]._element.getparent().remove(self.doc.paragraphs[0]._element)

        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        cell = table.cell(0, 0)
        width = Cm(15)
        cell.width = width
        tbl = table._tbl
        tblPr = tbl.tblPr
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:type"), "dxa")
        tblW.set(qn("w:w"), "8505")
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        set_cell_shading(cell, "F7FAFC")
        self._card_borders(cell)
        self._card_margins(cell)

        title = self.cover.get("title", "Docker")
        subtitle = self.cover.get("subtitle", "")
        name = self.cover.get("name", "Dasari Sindhu")
        meta = self.cover.get("meta", [])

        self._card_line(cell, title, 40, NAVY, bold=True, before=2, after=0, first=True)
        if subtitle:
            self._card_line(cell, subtitle, 13, RGBColor(0x3D, 0x51, 0x66), before=4, after=2)
        rule = cell.add_paragraph()
        rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rule.paragraph_format.space_before = Pt(8)
        rule.paragraph_format.space_after = Pt(8)
        rule.paragraph_format.left_indent = Cm(4.2)
        rule.paragraph_format.right_indent = Cm(4.2)
        self._rule(rule, "bottom", "7AA0C4", "12")
        self._card_line(cell, name, 18, INK, bold=True, before=6, after=2)
        for index, line in enumerate(meta):
            shown = line.replace(" | ", "   ·   ")
            last = index == len(meta) - 1
            self._card_line(
                cell,
                shown,
                11,
                RGBColor(0x3D, 0x51, 0x66),
                before=0,
                after=8 if last else 1,
            )

    def _card_line(self, cell, text, size, color, bold=False, before=0, after=0, first=False):
        paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.line_spacing = 1.05
        run = paragraph.add_run(text)
        set_run_font(run, "Calibri", size, color, bold=bold)
        return paragraph

    def _card_borders(self, cell):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge, color, size in (
            ("top", "1A3654", "16"),
            ("left", "D5E3EF", "8"),
            ("bottom", "D5E3EF", "8"),
            ("right", "D5E3EF", "8"),
        ):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), size)
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
            borders.append(el)
        tcPr.append(borders)

    def _card_margins(self, cell):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        mar = OxmlElement("w:tcMar")
        for edge, val in (("top", "280"), ("left", "360"), ("bottom", "280"), ("right", "360")):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:w"), val)
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tcPr.append(mar)

    def _render_table(self):
        rows = self.table_rows or []
        rows = [row for row in rows if row]
        if not rows:
            return
        cols = max(len(row) for row in rows)
        table = self.doc.add_table(rows=len(rows), cols=cols)
        table.autofit = False
        table.allow_autofit = False
        width = self.usable
        for row_index, row_data in enumerate(rows):
            row = table.rows[row_index]
            prevent_row_split(row)
            for col_index in range(cols):
                cell = row.cells[col_index]
                cell.width = int(width / cols)
                set_cell_borders(cell)
                set_cell_margins(cell)
                info = row_data[col_index] if col_index < len(row_data) else {"header": False, "runs": []}
                if info["header"] or row_index == 0:
                    set_cell_shading(cell, "EEF4F9")
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                color = NAVY if info["header"] or row_index == 0 else INK
                runs = info["runs"] or [("", set())]
                if info["header"] or row_index == 0:
                    runs = [(text, flags | {"bold"}) for text, flags in runs]
                self._add_runs(paragraph, runs, size=10, color=color)
        # Spacer after the table.
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(2)
        spacer.paragraph_format.space_after = Pt(6)

    def build(self, html):
        self.feed(html)
        self._close_capture()
        return self.doc


def main():
    html = HTML.read_text(encoding="utf-8")
    doc = Builder().build(html)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(OUT)
        print(OUT)
    except PermissionError:
        alt = OUT.with_name(OUT.stem + "_Print.docx")
        doc.save(alt)
        print(alt)


if __name__ == "__main__":
    main()
